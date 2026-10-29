import io
import base64
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from pypdf import PdfWriter
from evida.source_formats import pdf_pages, xlsx_rows
from evida.repository_documents import read_repository_document
from evida.runner import Runner
from evida.store import Store


class SourceFormats(unittest.TestCase):
    def test_embedded_image_is_reported_as_unread_original_not_missing_content(self):
        book = Workbook(); sheet = book.active
        sheet.append(['Candidate', 'Structure'])
        sheet.append(['C001', None])
        png = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aCuoAAAAASUVORK5CYII=')
        sheet.add_image(Image(io.BytesIO(png)), 'B2')
        data = io.BytesIO(); book.save(data)
        rows, summary, _ = xlsx_rows(data.getvalue(), 'https://example.invalid/structures.xlsx')
        self.assertIsNone(rows[1]['cells'][1]['value'])
        self.assertTrue(rows[1]['requires_original_review'])
        content = summary['non_cell_content']
        self.assertEqual(content['media_file_count'], 1)
        self.assertFalse(content['rendered_here'])

    def test_spreadsheet_keeps_zero_blank_literal_formula_and_locations(self):
        book = Workbook()
        sheet = book.active
        sheet.title = 'Context'
        sheet.append(['Candidate', 'Observed', 'Missing', 'Calculation'])
        sheet.append(['C001', 0, None, '=B2+1'])
        sheet.merge_cells('A3:B3')
        sheet['A3'] = 'Header spans two columns'
        sheet.row_dimensions[2].hidden = True
        sheet.column_dimensions['D'].hidden = True
        data = io.BytesIO(); book.save(data)
        rows, summary, _ = xlsx_rows(data.getvalue(), 'https://example.invalid/data.xlsx')
        self.assertEqual(rows[1]['cells'][1]['value'], 0)
        self.assertIsNone(rows[1]['cells'][2]['value'])
        self.assertEqual(rows[1]['cells'][3]['value'], '=B2+1')
        self.assertIsNone(rows[1]['cells'][3]['cached_value'])
        self.assertFalse(rows[1]['cells'][3]['formula_evaluated_here'])
        self.assertEqual(rows[1]['cells'][3]['coordinate'], 'D2')
        self.assertTrue(rows[1]['hidden_row'])
        self.assertEqual(summary['sheets'][0]['merged_ranges'], ['A3:B3'])
        self.assertEqual(summary['sheets'][0]['hidden_columns'], ['D'])

    def test_empty_pdf_pages_survive_as_unread_not_absent_evidence(self):
        writer = PdfWriter(); writer.add_blank_page(width=612, height=792)
        raw = io.BytesIO(); writer.write(raw)
        rows, summary, _ = pdf_pages(raw.getvalue(), 'https://example.invalid/pages.pdf')
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['page_number'], 1)
        self.assertTrue(rows[0]['text_layer_empty'])
        self.assertTrue(rows[0]['requires_original_review'])
        self.assertTrue(rows[0]['source_url'].endswith('#page=1'))
        self.assertEqual(summary['empty_text_pages'], 1)

    def test_pdf_original_is_stored_but_not_sent_in_reading(self):
        writer = PdfWriter(); writer.add_blank_page(width=612, height=792)
        raw = io.BytesIO(); writer.write(raw)
        record = {'id': 123, 'files': [{'key': 'source.pdf'}]}
        with patch('evida.repository_documents.fetch', side_effect=[(json.dumps(record).encode(), {}), (raw.getvalue(), {})]):
            value = read_repository_document({'record_id':'123','filename':'source.pdf'})
        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder)/'db.sqlite3'); wid = store.create('fixture')['id']
            aid = store.add_artifact(wid, 'pdf', 'repository_document', json.dumps(value).encode(), {})
            runner = object.__new__(Runner); runner.store = store
            view = runner.inspect(wid, aid)
            self.assertIn('original_file_base64', value)
            self.assertNotIn('original_file_base64', view['result'])
            self.assertNotIn('metadata_response', view['result'])
            self.assertEqual(view['result']['total_rows'], 1)


if __name__ == '__main__':
    unittest.main()
