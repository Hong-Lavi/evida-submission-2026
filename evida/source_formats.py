"""Source locations and literal values, without inferring scientific relationships."""
import datetime
import importlib.metadata
import io
import zipfile


def pdf_pages(raw, source_url):
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(raw))
    if reader.is_encrypted:
        raise ValueError('암호화 PDF는 여기서 읽지 않습니다. 접근 가능한 원본이 필요합니다.')
    if len(reader.pages) > 500:
        raise ValueError('이 PDF는 현재 문서 처리 범위를 넘습니다. 필요한 원문 구간을 확인해 주세요.')
    rows = []
    for number, page in enumerate(reader.pages, 1):
        text = '' if page.get_contents() is None else page.extract_text(extraction_mode='layout', layout_mode_space_vertically=False)
        rows.append({'row_id': f'pdf:page:{number}', 'kind': 'pdf_page', 'page_number': number,
                     'text': text, 'source_url': source_url + f'#page={number}',
                     'requires_original_review': True, 'text_layer_empty': not bool(text.strip()),
                     'page_size_points': [float(page.mediabox.width), float(page.mediabox.height)],
                     'rotation_degrees': page.rotation,
                     'extraction': 'layout text; spaces and ordering are not verified panel associations'})
    return rows, {'pages': len(rows), 'empty_text_pages': sum(r['text_layer_empty'] for r in rows)}, {
        'pypdf': importlib.metadata.version('pypdf')}


def _cell_value(value):
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, datetime.timedelta):
        return {'duration_seconds': value.total_seconds()}
    return value


def xlsx_rows(raw, source_url):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
            raise ValueError('엑셀의 압축 해제 크기가 현재 처리 범위를 넘습니다.')
        non_cell_parts = [i.filename for i in archive.infolist()
                          if i.filename.startswith(('xl/drawings/', 'xl/media/', 'xl/embeddings/'))
                          and not i.filename.endswith(('/', '.rels'))]
    book = load_workbook(io.BytesIO(raw), data_only=False, keep_links=False)
    cached = load_workbook(io.BytesIO(raw), data_only=True, keep_links=False)
    rows, sheets = [], []
    grid_cells = 0
    try:
        for sheet in book:
            grid_cells += sheet.max_row * sheet.max_column
            # The official CACHE structure workbook contains 393,954 grid cells.
            # Keep a finite limit while allowing that measured, nontruncated input.
            if grid_cells > 1000000:
                raise ValueError('엑셀 셀 범위가 현재 처리 범위를 넘습니다. 필요한 시트·범위 확인이 필요합니다.')
            merged = [str(r) for r in sheet.merged_cells.ranges]
            hidden_columns = [k for k, v in sheet.column_dimensions.items() if v.hidden]
            sheets.append({'name': sheet.title, 'state': sheet.sheet_state,
                           'dimensions': sheet.calculate_dimension(), 'merged_ranges': merged,
                           'hidden_columns': hidden_columns})
            for number, values in enumerate(sheet.iter_rows(), 1):
                cells = []
                for cell in values:
                    coordinate = f'{get_column_letter(cell.column)}{number}'
                    entry = {'coordinate': coordinate, 'value': _cell_value(cell.value),
                             'data_type': cell.data_type, 'number_format': cell.number_format}
                    if cell.data_type == 'f':
                        entry['cached_value'] = _cell_value(cached[sheet.title][coordinate].value)
                        entry['formula_evaluated_here'] = False
                    if cell.comment:
                        entry['comment'] = cell.comment.text
                    cells.append(entry)
                if not any(c['value'] is not None or 'comment' in c for c in cells):
                    continue
                rows.append({'row_id': f'xlsx:{sheet.title}:row:{number}', 'kind': 'xlsx_row',
                             'sheet': sheet.title, 'row_number': number, 'cells': cells,
                             'hidden_row': bool(sheet.row_dimensions[number].hidden),
                             'source_url': source_url, 'requires_original_review': bool(merged or non_cell_parts),
                             'text': ' | '.join(f"{c['coordinate']}={c['value']!r}" for c in cells)})
    finally:
        book.close()
        cached.close()
    return rows, {'sheets': sheets, 'nonempty_rows': len(rows), 'grid_cells': grid_cells,
                  'non_cell_content': {'package_part_count': len(non_cell_parts),
                                       'media_file_count': sum(n.startswith('xl/media/') for n in non_cell_parts),
                                       'drawing_file_count': sum(n.startswith('xl/drawings/') for n in non_cell_parts),
                                       'embedded_object_file_count': sum(n.startswith('xl/embeddings/') for n in non_cell_parts),
                                       'rendered_here': False,
                                       'interpretation': '그림·차트·삽입 개체는 원본에 보존됩니다. 셀 값만 읽어 그 내용이나 과학적 대응이 없다고 판단할 수 없습니다.'}}, {
        'openpyxl': importlib.metadata.version('openpyxl')}
